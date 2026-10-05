from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADERS = ["File Name", "Pg. No", "Type", "Author", "Modified Time", "Comments"]


def write_safe_cell(sheet, row, column, value):
    """Write all extracted strings as literal XLSX text, never formulas.

    Strip XML-prohibited control characters. Preserve Unicode, whitespace,
    and formula-looking text without adding a visible apostrophe.
    Numeric IDs/page numbers and missing values retain their original types.
    """
    cell = sheet.cell(row=row, column=column)
    if isinstance(value, str):
        cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
        # Assignment can infer 'f' for '=' or 'e' for Excel error strings.
        # Override it explicitly so the XLSX writer emits a text cell.
        cell.data_type = "s"
        cell.number_format = "@"
    else:
        cell.value = value
    return cell


def generate_excel(records, output_path, filename=""):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Annotations"
    for col, header in enumerate(HEADERS, start=1):
        cell = write_safe_cell(sheet, 1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_idx, record in enumerate(records, start=2):
        values = [filename or record.get("filename", ""), record.get("page_number"), record.get("annotation_type") or record.get("object_type", ""), record.get("author", ""), record.get("modified_time", ""), record.get("text", "")]
        for col_idx, value in enumerate(values, start=1):
            write_safe_cell(sheet, row_idx, col_idx, value)
    widths = [35, 10, 20, 25, 20, 75]
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    sheet.freeze_panes = "A2"
    workbook.save(output_path)
